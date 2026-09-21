# -*- coding: utf-8 -*-
"""
生成《作业1-图像处理》Word报告文档
根据课堂要求（形态学变换、高斯滤波、高斯金字塔）整理 main.py 中的代码，
分块写入 docx 文档，并简单说明实现思路与运行效果。
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


OUTPUT_PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                            "作业1-图像处理报告.docx")


# ===================== 从 main.py 中整理出的三块核心代码 =====================
MORPHOLOGY_CODE = '''# ===================== 1. 形态学变换 =====================
def morphology_demo(img_path):
    # 读取灰度图像
    img_gray = cv2.imread(img_path, cv2.IMREAD_GRAYSCALE)
    # 二值化
    _, img_bin = cv2.threshold(img_gray, 127, 255, cv2.THRESH_BINARY)

    # 自定义3*3矩形结构化元素
    kernel = np.ones((3, 3), dtype=np.uint8)

    # 腐蚀、膨胀、开运算、闭运算
    img_erode = cv2.erode(img_bin, kernel)
    img_dilate = cv2.dilate(img_bin, kernel)
    img_open = cv2.morphologyEx(img_bin, cv2.MORPH_OPEN, kernel)
    img_close = cv2.morphologyEx(img_bin, cv2.MORPH_CLOSE, kernel)

    # 可视化
    titles = ["原始灰度", "二值图", "腐蚀", "膨胀", "开运算", "闭运算"]
    imgs = [img_gray, img_bin, img_erode, img_dilate, img_open, img_close]
    plt.figure(figsize=(12, 7))
    for i in range(6):
        plt.subplot(2, 3, i + 1)
        plt.imshow(imgs[i], cmap="gray")
        plt.title(titles[i])
        plt.axis("off")
    plt.suptitle("形态学变换结果")
    plt.tight_layout()
    plt.show()
    return img_bin
'''

GAUSSIAN_FILTER_CODE = '''# ===================== 2. 手写高斯滤波函数 gaussian_filter =====================
def gaussian_filter(src, kernel_size=5, sigma=1.5):
    """
    手写高斯滤波，不调用cv2.GaussianBlur
    :param src:输入灰度图
    :param kernel_size:核大小，奇数
    :param sigma:高斯标准差
    :return:滤波后图像
    """
    # 生成高斯核
    k = kernel_size // 2
    x, y = np.meshgrid(np.arange(-k, k+1), np.arange(-k, k+1))
    g = np.exp(-(x**2 + y**2)/(2*sigma**2))
    g = g / np.sum(g)  # 归一化

    # 边界填充
    src_pad = np.pad(src, pad_width=k, mode="reflect")
    out = np.zeros_like(src, dtype=np.float32)

    # 卷积运算
    h, w = src.shape
    for i in range(h):
        for j in range(w):
            patch = src_pad[i:i+kernel_size, j:j+kernel_size]
            out[i, j] = np.sum(patch * g)
    out = np.clip(out, 0, 255).astype(np.uint8)
    return out
'''

GAUSSIAN_PYRAMID_CODE = '''# ===================== 3. 高斯金字塔（基于手写gaussian_filter） =====================
def gaussian_pyramid(src, level=4, kernel_size=5, sigma=1.5):
    """
    生成高斯金字塔
    :param src:输入灰度图像
    :param level:金字塔层数
    :return:金字塔列表，第0层原图
    """
    pyramid = []
    current = src.copy()
    pyramid.append(current)
    for _ in range(level-1):
        # 先高斯滤波
        blur = gaussian_filter(current, kernel_size, sigma)
        # 下采样：隔点删除
        next_level = blur[::2, ::2]
        pyramid.append(next_level)
        current = next_level
    # 可视化金字塔
    plt.figure(figsize=(14, 4))
    for idx, layer in enumerate(pyramid):
        plt.subplot(1, level, idx+1)
        plt.imshow(layer, cmap="gray")
        plt.title(f"Level {idx}, shape{layer.shape}")
        plt.axis("off")
    plt.suptitle("高斯金字塔")
    plt.tight_layout()
    plt.show()
    return pyramid
'''

MAIN_CODE = '''import cv2
import numpy as np
import matplotlib.pyplot as plt

# 解决中文显示为方块/乱码的问题
plt.rcParams['font.sans-serif'] = ['Microsoft YaHei']
# 解决负号显示异常的问题
plt.rcParams['axes.unicode_minus'] = False

if __name__ == "__main__":
    image_file = "blobs.png"  # 修改成你的图片路径

    # 任务1 形态学变换
    binary_img = morphology_demo(image_file)

    # 任务2 高斯滤波测试
    blur_result = gaussian_filter(binary_img, kernel_size=5, sigma=2.0)
    plt.figure()
    plt.subplot(121), plt.imshow(binary_img, cmap="gray"), plt.title("原图"), plt.axis("off")
    plt.subplot(122), plt.imshow(blur_result, cmap="gray"), plt.title("高斯滤波"), plt.axis("off")
    plt.show()

    # 任务3 高斯金字塔
    gaussian_pyramid(binary_img, level=4, kernel_size=5, sigma=2.0)
'''


def set_cn_font(run, name="微软雅黑", size=12, bold=False, color=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = name
    # 设置中文字体（东亚字体）
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_heading(doc, text, level=1, size=16, color=(0x2E, 0x74, 0xB5)):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_cn_font(run, name="微软雅黑", size=size, bold=True, color=color)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_body(doc, text, size=12, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_cn_font(run, name="微软雅黑", size=size, bold=bold)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.3
    return p


def add_code_block(doc, code_text):
    """以等宽字体、带灰色底纹的方式插入代码块"""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.left_indent = Cm(0.3)
    # 段落底纹（浅灰色背景）
    pPr = p._p.get_or_add_pPr()
    shd = pPr.makeelement(qn('w:shd'), {
        qn('w:val'): 'clear',
        qn('w:color'): 'auto',
        qn('w:fill'): 'F2F2F2',
    })
    pPr.append(shd)

    lines = code_text.strip("\n").split("\n")
    for idx, line in enumerate(lines):
        run = p.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(10)
        run._element.rPr.rFonts.set(qn('w:eastAsia'), "Consolas")
        if idx != len(lines) - 1:
            run.add_break()
    return p


def add_image_if_exists(doc, path, width_cm=14, caption=None):
    if os.path.exists(path):
        doc.add_picture(path, width=Cm(width_cm))
        last_p = doc.paragraphs[-1]
        last_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if caption:
            cap = doc.add_paragraph()
            run = cap.add_run(caption)
            set_cn_font(run, size=10)
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER


def build_document():
    doc = Document()

    # 设置正文默认字体
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

    # ========== 标题 ==========
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("作业1 - 图像处理")
    set_cn_font(title_run, name="微软雅黑", size=22, bold=True, color=(0xC0, 0x50, 0x4D))
    title_p.paragraph_format.space_after = Pt(6)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run("基于 OpenCV 的形态学变换、手写高斯滤波与高斯金字塔实现")
    set_cn_font(sub_run, size=12)
    sub_p.paragraph_format.space_after = Pt(18)

    doc.add_paragraph().paragraph_format.space_after = Pt(0)

    # ========== 一、形态学变换 ==========
    add_heading(doc, "一、形态学变换", size=16)
    add_body(
        doc,
        "任务要求：使用 OpenCV，读取图像为灰度并二值化（可使用 OpenCV 函数），"
        "自定义结构化元素（如 3*3），实现膨胀、腐蚀、开、闭等函数，并可视化结果。\n"
        "参考资料：https://www.mathworks.com/help/images/morphological-dilation-and-erosion.html"
    )
    add_body(doc, "实现代码：", bold=True)
    add_code_block(doc, MORPHOLOGY_CODE)
    add_body(
        doc,
        "实现说明：先用 cv2.imread 以灰度模式读取图像，再用 cv2.threshold 进行二值化；"
        "使用 np.ones((3,3), dtype=np.uint8) 自定义 3×3 结构化元素；分别调用 cv2.erode、"
        "cv2.dilate、cv2.morphologyEx(MORPH_OPEN)、cv2.morphologyEx(MORPH_CLOSE) 完成腐蚀、"
        "膨胀、开运算、闭运算，并用 matplotlib 以 2×3 子图的形式可视化原图、二值图及四种形态学"
        "处理结果。"
    )
    add_image_if_exists(doc, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                                            "results", "morphology_result.png"),
                         caption="图1 形态学变换结果示意图（如已生成运行截图，将自动插入此处）")

    # ========== 二、高斯滤波 ==========
    add_heading(doc, "二、高斯滤波", size=16)
    add_body(
        doc,
        "任务要求：写一个高斯 filter 函数，对图像执行高斯滤波（不调用 cv2.GaussianBlur）。"
    )
    add_body(doc, "实现代码：", bold=True)
    add_code_block(doc, GAUSSIAN_FILTER_CODE)
    add_body(
        doc,
        "实现说明：首先根据 kernel_size 和 sigma 生成二维高斯核 g，并对其归一化，保证滤波前后"
        "图像整体亮度不变；对原图使用 reflect 方式做边界填充，避免边缘信息丢失；再通过双重循环"
        "对每个像素位置取出对应大小的图像块（patch），与高斯核逐元素相乘后求和，得到该像素的滤波"
        "结果；最后将结果裁剪到 [0,255] 并转换为 uint8 类型输出。"
    )

    # ========== 三、高斯金字塔 ==========
    add_heading(doc, "三、高斯金字塔", size=16)
    add_body(
        doc,
        "任务要求：基于高斯 filter，生成高斯金字塔。"
    )
    add_body(doc, "实现代码：", bold=True)
    add_code_block(doc, GAUSSIAN_PYRAMID_CODE)
    add_body(
        doc,
        "实现说明：以第一步二值化后的图像作为金字塔第0层；每一层先调用自定义的 gaussian_filter "
        "对当前层做高斯平滑，再通过隔点采样（[::2, ::2]）将图像长宽各缩小一半，得到下一层图像，"
        "循环 level-1 次后得到完整的高斯金字塔；最后用 matplotlib 将各层图像并排显示，标题中标注"
        "每层的层号与图像尺寸。"
    )

    # ========== 四、主程序调用 ==========
    add_heading(doc, "四、主程序（main.py）调用流程", size=16)
    add_code_block(doc, MAIN_CODE)
    add_body(
        doc,
        "调用说明：程序入口先读取图像 blobs.png，依次执行任务1（形态学变换，得到二值图 "
        "binary_img）、任务2（对 binary_img 做高斯滤波，并对比原图与滤波结果）、任务3（基于"
        "binary_img 生成4层高斯金字塔并可视化）。同时在文件开头设置了 matplotlib 的中文字体"
        "（Microsoft YaHei）及负号显示，解决了中文标题/坐标轴显示为方块乱码的问题。"
    )

    # ========== 五、作业提交说明 ==========
    add_heading(doc, "五、作业提交说明", size=16)
    add_body(doc, "作业报告：Word 文档（即本文档）。")
    add_body(doc, "截止时间：9月27日21点前发给助教。")

    return doc


if __name__ == "__main__":
    document = build_document()
    document.save(OUTPUT_PATH)
    print("文档已生成:", OUTPUT_PATH)
